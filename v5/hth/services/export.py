"""Excel export of any tabular report."""

from __future__ import annotations

import datetime as dt
import re
from collections.abc import Sequence
from pathlib import Path


def safe_filename(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", name).strip("_") or "report"


def export_rows(
    path: Path,
    columns: Sequence[str],
    rows: Sequence[Sequence[object]],
    title: str = "Report",
) -> Path:
    # openpyxl (and the numpy it pulls in) takes about a second to load: only when exporting
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = safe_filename(title)[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Generated {dt.datetime.now():%Y-%m-%d %H:%M}"])
    ws.append([])
    ws.append(list(columns))
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E4F4F")
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(["" if v is None else v for v in row])
    for idx, col in enumerate(columns, start=1):
        width = max([len(str(col))] + [len(str(r[idx - 1])) for r in rows if idx - 1 < len(r)] or [8])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 8), 60)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
